import streamlit as st
import pandas as pd
import re
import io
from openpyxl.styles import Font as ExcelFont
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, Font as DrawFont
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.colors import ColorChoice

# 人間と同じように数字の大きさを認識して並び替える関数
def natural_sort_key(uploaded_file):
    return [int(text) if text.isdigit() else text.lower() for text in re.split(r'(\d+)', uploaded_file.name)]

def main():
    st.title("🧪 ESR Data Master")
    st.write("テキストファイルの読み込み・Excel変換から、グラフ作成までを実行します。")

    # 1. テキストファイルのアップロード
    uploaded_files = st.file_uploader(
        "フォルダ内のテキストファイル（.txt）をすべて選択してください", 
        type="txt", 
        accept_multiple_files=True
    )

    if uploaded_files:
        # ファイルを自然順（5min -> 30minなど）に並び替える
        uploaded_files = sorted(uploaded_files, key=natural_sort_key)
        
        # ---------------------------------------------------------
        # 【事前準備】アップロードされたファイル名から「作成されるシート名」を予測する
        # （Excelを作る前にUIを表示するため）
        # ---------------------------------------------------------
        file_to_sheet = {}
        ordered_sheet_names = []
        seen_sheet_names = set()
        
        for f in uploaded_files:
            raw_name = re.sub(r'\.txt$', '', f.name, flags=re.IGNORECASE)
            
            # 31文字制限カット（真ん中省略）
            if len(raw_name) > 31:
                sheet_name = raw_name[:15] + ".." + raw_name[-14:]
            else:
                sheet_name = raw_name
            
            # 重複回避処理
            original_sheet_name = sheet_name
            counter = 2
            while sheet_name in seen_sheet_names:
                suffix = f"_{counter}"
                sheet_name = original_sheet_name[:31 - len(suffix)] + suffix
                counter += 1
                
            seen_sheet_names.add(sheet_name)
            ordered_sheet_names.append(sheet_name)
            file_to_sheet[f.name] = sheet_name
        
        st.info(f"{len(uploaded_files)} 個のファイルが読み込まれました。")
        
        # ---------------------------------------------------------
        # 2. UI：グラフを作成するシートとマーカーの選択
        # ---------------------------------------------------------
        st.markdown("### ⚙️ グラフ作成の設定")
        st.write("グラフを作成したいデータと、そのX軸計算に使うMnマーカーの組み合わせを選んでください。")
        
        target_sheet_names = st.multiselect(
            "グラフを作成するシートを選択してください（複数選択可）", 
            ordered_sheet_names, 
            default=[ordered_sheet_names[0]] if ordered_sheet_names else None
        )

        sheet_mapping = {}
        if target_sheet_names:
            for target in target_sheet_names:
                # デフォルトで一番最後のシート（大抵はMnマーカー）が選ばれるようにする
                mn_sheet = st.selectbox(
                    f"「{target}」のMnマーカー", 
                    ordered_sheet_names, 
                    index=len(ordered_sheet_names)-1 if ordered_sheet_names else 0,
                    key=f"mn_{target}"
                )
                sheet_mapping[target] = mn_sheet

        # ---------------------------------------------------------
        # 3. 実行ボタンが押された後の処理
        # ---------------------------------------------------------
        if st.button("変換とグラフ作成を一括実行する"):
            excel_buffer = io.BytesIO()
            progress_bar = st.progress(0)
            
            yu_font = ExcelFont(name='游ゴシック')
            
            with pd.ExcelWriter(excel_buffer, engine='openpyxl') as writer:
                
                # =====================================================
                # フェーズ1：テキストをExcelデータに変換
                # =====================================================
                st.text("テキストデータをExcelに変換中...")
                for i, uploaded_file in enumerate(uploaded_files):
                    content = uploaded_file.getvalue().decode('utf-8', errors='ignore')
                    
                    parsed_data = []
                    for line in content.splitlines():
                        line = line.rstrip() # 先頭の空白を残す（A列を空けるため）
                        if not line:
                            continue
                        
                        row = re.split(r'[ ,\t]+', line)
                        
                        clean_row = []
                        for cell in row:
                            if cell.startswith('='):
                                clean_row.append("'" + cell)
                            else:
                                try:
                                    if '.' in cell:
                                        clean_row.append(float(cell))
                                    else:
                                        clean_row.append(int(cell))
                                except ValueError:
                                    clean_row.append(cell)
                                    
                        parsed_data.append(clean_row)
                    
                    df = pd.DataFrame(parsed_data)
                    sheet_name = file_to_sheet[uploaded_file.name]
                    
                    # Excelに書き込み
                    df.to_excel(writer, sheet_name=sheet_name, index=False, header=False)
                    
                    # 游ゴシック適用
                    worksheet = writer.sheets[sheet_name]
                    for row in worksheet.iter_rows():
                        for cell in row:
                            cell.font = yu_font
                            
                    # プログレスバー（前半50%を使用）
                    progress_bar.progress((i + 1) / (len(uploaded_files) * 2))

                # =====================================================
                # フェーズ2：作成したシートにグラフと計算式を追加
                # =====================================================
                st.text("グラフを作成中...")
                total_graphs = len(sheet_mapping)
                
                for idx, (target_sheet, marker_sheet) in enumerate(sheet_mapping.items()):
                    ws_target = writer.sheets[target_sheet]
                    safe_marker_sheet = f"'{marker_sheet}'"
                    
                    # ① X軸の計算式
                    ws_target['A80'] = f"={safe_marker_sheet}!$D$12"
                    step_val_formula = f"({safe_marker_sheet}!$D$13-{safe_marker_sheet}!$D$12)/4095"
                    
                    for i in range(1, 4096):
                        row = 80 + i
                        ws_target[f'A{row}'] = f"=$A$80+(ROW()-80)*{step_val_formula}"

                    # ② グラフの生成とベース設定
                    chart = ScatterChart()
                    chart.varyColors = False
                    chart.title = None
                    chart.legend = None  
                    chart.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))

                    xvalues = Reference(ws_target, min_col=1, min_row=80, max_row=4175)
                    yvalues = Reference(ws_target, min_col=2, min_row=80, max_row=4175)
                    
                    ser = Series(yvalues, xvalues, title_from_data=False)
                    ser.smooth = False
                    ser.marker = Marker(symbol="none")
                    ser.graphicalProperties.line.solidFill = "000000"
                    ser.graphicalProperties.line.width = 19050  
                    chart.series.append(ser)

                    # ③ 軸の表示設定（完全下部・左側固定）
                    chart.x_axis.delete = False
                    chart.y_axis.delete = False
                    chart.x_axis.crosses = "min"
                    chart.y_axis.crosses = "min"
                    chart.x_axis.tickLblPos = "low"
                    chart.y_axis.tickLblPos = "low"
                    chart.x_axis.title = None
                    chart.y_axis.title = None
                    chart.x_axis.number_format = 'General'
                    chart.y_axis.number_format = 'General'
                    chart.x_axis.majorGridlines = None
                    chart.y_axis.majorGridlines = None
                    chart.x_axis.majorTickMark = "out"
                    chart.y_axis.majorTickMark = "out"

                    axis_line = LineProperties(solidFill=ColorChoice(prstClr="black"), w=9525)
                    chart.x_axis.spPr = GraphicalProperties(ln=axis_line)
                    chart.y_axis.spPr = GraphicalProperties(ln=axis_line)

                    # ④ フォント設定 (Times New Roman)
                    times_font = DrawFont(typeface='Times New Roman')
                    cp = CharacterProperties(latin=times_font, sz=1100)
                    pPr = ParagraphProperties(defRPr=cp)
                    
                    chart.x_axis.txPr = RichText(p=[Paragraph(pPr=pPr, endParaRPr=cp)])
                    chart.y_axis.txPr = RichText(p=[Paragraph(pPr=pPr, endParaRPr=cp)])

                    # ⑤ 配置
                    ws_target.add_chart(chart, "H3")

                    # プログレスバー（後半50%を使用）
                    progress_bar.progress(0.5 + ((idx + 1) / total_graphs) * 0.5)

            # =====================================================
            # 完了とダウンロード
            # =====================================================
            st.success("🎉 データ変換とグラフ作成がすべて完了しました！")
            
            st.download_button(
                label="📁 完成したExcelファイルをダウンロード",
                data=excel_buffer.getvalue(),
                file_name="ESR_Combined_Graph_Data.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

if __name__ == "__main__":
    main()